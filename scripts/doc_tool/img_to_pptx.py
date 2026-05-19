import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

from scripts.common_py.log import guohub_error_print, guohub_json_print, guohub_logger


def natural_sort_key(s: str):
    return [int(c) if c.isdigit() else c.lower() for c in re.split(r"(\d+)", s)]


def main():
    parser = argparse.ArgumentParser(description="将图片合成为 PPTX（每页一张图）")
    parser.add_argument("input_dir", help="图片所在目录")
    parser.add_argument("-o", "--output", default=None, help="输出 PPTX 路径（默认在输入目录下生成）")
    parser.add_argument("--ext", nargs="*", default=[".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp"],
                        help="图片扩展名过滤（默认常见图片格式）")
    args = parser.parse_args()

    input_dir = Path(args.input_dir)
    if not input_dir.is_dir():
        guohub_error_print(f"目录不存在: {input_dir}")

    exts = set(args.ext)
    files = sorted(
        [f for f in input_dir.iterdir() if f.is_file() and f.suffix.lower() in exts],
        key=lambda f: natural_sort_key(f.stem),
    )

    if not files:
        guohub_error_print(f"目录中没有找到图片: {input_dir}")

    guohub_logger.info(f"找到 {len(files)} 张图片")

    from PIL import Image
    first = Image.open(files[0])
    w, h = first.size
    first.close()

    prs = Presentation()
    prs.slide_width = Emu(w * 914400 // 96)
    prs.slide_height = Emu(h * 914400 // 96)

    for f in files:
        slide = prs.slides.add_slide(next((l for l in prs.slide_layouts if "blank" in l.name.lower()), prs.slide_layouts[0]))
        slide.shapes.add_picture(str(f), Emu(0), Emu(0), prs.slide_width, prs.slide_height)

    out_path = args.output or str(input_dir / f"{input_dir.name}.pptx")
    prs.save(out_path)
    guohub_json_print({"output": str(out_path), "pages": len(files), "width": w, "height": h})


if __name__ == "__main__":
    main()
