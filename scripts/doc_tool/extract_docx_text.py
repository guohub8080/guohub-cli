import argparse
from pathlib import Path

try:
    from docx import Document
except ImportError:
    import sys
    sys.exit("缺少 python-docx，请运行: pnpm dev doctor")

from scripts.common_py.log import guohub_logger, guohub_error_print


def extract_text(docx_path: Path, output_path: Path | None = None) -> str:
    doc = Document(str(docx_path))

    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    full_text = "\n\n".join(paragraphs)

    if output_path:
        output_path.write_text(full_text, encoding="utf-8")
        guohub_logger.info(f"文字已保存: {output_path}")

    return full_text


def main():
    parser = argparse.ArgumentParser(description="从 Word 文档中提取文字")
    parser.add_argument("file", help="Word 文档路径（.docx）")
    parser.add_argument("--output", "-o", default=None, help="输出文件路径（默认输出到控制台）")
    args = parser.parse_args()

    path = Path(args.file).expanduser().resolve()
    if not path.exists():
        guohub_error_print(f"文件不存在: {path}")
        return

    if path.suffix.lower() != ".docx":
        guohub_error_print(f"仅支持 .docx 格式，当前: {path.suffix}")
        return

    output = None
    if args.output:
        output = Path(args.output).expanduser().resolve()

    text = extract_text(path, output)

    if not args.output:
        print(text)


if __name__ == "__main__":
    main()
